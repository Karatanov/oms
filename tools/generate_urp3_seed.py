"""Generate the canonical URP III seed from the supplied monitoring workbooks.

The production application never reads user-local Excel paths. This utility is
only the auditable transformation step; its normalized JSON and SQL outputs are
checked into the repository.
"""

from __future__ import annotations

import argparse
import json
import re
import uuid
from collections import OrderedDict, defaultdict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


NAMESPACE = uuid.UUID("cdad0bed-64ed-4a0b-a96f-67fdd4a1d308")
PROGRAMME_UUID = str(uuid.uuid5(NAMESPACE, "programme:urp-iii"))
SNAPSHOT_DATE = date(2026, 4, 24)
PROJECTED_UAH_PER_EUR = Decimal("48")


def cleaned(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = " ".join(value.replace("\r", " ").replace("\n", " ").split())
        if not value or value.startswith("#"):
            return None
        return value
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return Decimal(str(value))
    return value


def decimal_value(value) -> Decimal | None:
    value = cleaned(value)
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, str):
        candidate = value.replace(" ", "").replace(",", ".")
        try:
            return Decimal(candidate)
        except Exception:
            return None
    return None


def positive_decimal(value) -> Decimal | None:
    value = decimal_value(value)
    return value if value is not None and value > 0 else None


def sql(value) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''") + "'"


def json_value(value):
    if isinstance(value, Decimal):
        return format(value, "f")
    return value


def normalized_sector(value: str | None) -> str | None:
    text = (value or "").lower()
    for needle, code in (
        ("water supply", "water_supply"),
        ("sewerage", "sewerage"),
        ("heat supply", "heat_supply"),
        ("health", "health"),
        ("housing", "housing"),
        ("education", "education"),
    ):
        if needle in text:
            return code
    return None


def normalized_construction_type(value: str | None) -> str | None:
    text = (value or "").lower()
    if "new construction" in text:
        return "new_construction"
    if "reconstruction" in text:
        return "reconstruction"
    if "overhaul" in text:
        return "capital_repair"
    return None


def normalized_project_status(work_status: str | None) -> str:
    text = (work_status or "").lower()
    if "suspended" in text or "призупинено" in text:
        return "suspended"
    if "completed" in text or "заверш" in text:
        return "completed"
    if "progress" in text or "реалізації" in text:
        return "active"
    return "planned"


def workbook_rows(path: Path) -> list[dict]:
    formula_sheet = load_workbook(path, data_only=False, read_only=False)["DB Table"]
    value_sheet = load_workbook(path, data_only=True, read_only=False)["DB Table"]
    rows = []
    for row_number in range(16, 152):
        row = {"sourceRow": row_number}
        for column in range(1, 129):
            letter = get_column_letter(column)
            value = value_sheet.cell(row_number, column).value
            if value is None:
                raw = formula_sheet.cell(row_number, column).value
                value = None if isinstance(raw, str) and raw.startswith("=") else raw
            row[letter] = cleaned(value)
        rows.append(row)
    return rows


def merge_duplicate_lots(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    grouped: OrderedDict[str, list[dict]] = OrderedDict()
    for row in rows:
        grouped.setdefault(str(row["F"]), []).append(row)
    merged = []
    decisions = []
    master_columns = [get_column_letter(index) for index in range(2, 25)]
    later_columns = [get_column_letter(index) for index in range(25, 129)]
    for lot_id, candidates in grouped.items():
        result = {"sourceRows": [item["sourceRow"] for item in candidates]}
        for column in master_columns:
            result[column] = next((item[column] for item in candidates if item.get(column) is not None), None)
        for column in later_columns:
            result[column] = next((item[column] for item in reversed(candidates) if item.get(column) is not None), None)
        result["uuid"] = str(uuid.uuid5(NAMESPACE, "subproject:" + lot_id))
        result["sector"] = normalized_sector(result.get("O"))
        result["constructionType"] = normalized_construction_type(result.get("P"))
        result["status"] = normalized_project_status(result.get("BK"))
        merged.append(result)
        if len(candidates) > 1:
            decisions.append({
                "lotId": lot_id,
                "sourceRows": result["sourceRows"],
                "rule": "master data from the first non-empty cell; financial, procurement and delivery state from the last non-empty cell",
            })
    return merged, decisions


def supplement_contracts(path: Path) -> dict[str, dict[str, dict]]:
    result: dict[str, dict[str, dict]] = defaultdict(dict)
    workbook = load_workbook(path, data_only=True, read_only=False)
    for sheet in workbook.worksheets:
        for row in range(10, sheet.max_row + 1):
            code = cleaned(sheet.cell(row, 9).value)
            kind = cleaned(sheet.cell(row, 12).value)
            procurement_id = cleaned(sheet.cell(row, 13).value)
            if not code or kind not in {"W", "TS", "CSC"} or not procurement_id:
                continue
            result[str(code)][str(kind)] = {
                "sourceSheet": sheet.title,
                "procurementId": procurement_id,
                "procurementMethod": cleaned(sheet.cell(row, 23).value),
                "tenderDocumentType": cleaned(sheet.cell(row, 24).value),
                "publishedInOjeu": cleaned(sheet.cell(row, 25).value),
                "estimatedProzorroDate": cleaned(sheet.cell(row, 26).value),
                "estimatedBidSubmissionDate": cleaned(sheet.cell(row, 27).value),
                "estimatedContractDate": cleaned(sheet.cell(row, 28).value),
                "estimatedContractEndDate": cleaned(sheet.cell(row, 29).value),
                "procurementStatus": cleaned(sheet.cell(row, 30).value),
                "comments": cleaned(sheet.cell(row, 33).value),
            }
    return result


def amount_record(kind: str, uah, eur) -> dict | None:
    uah_value = positive_decimal(uah)
    eur_value = positive_decimal(eur)
    if uah_value is None and eur_value is None:
        return None
    if uah_value is None:
        uah_value = eur_value * PROJECTED_UAH_PER_EUR
    if eur_value is None:
        eur_value = uah_value / PROJECTED_UAH_PER_EUR
    return {
        "kind": kind,
        "amount": uah_value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
        "currency": "UAH",
        "convertedAmount": eur_value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
        "uahPerEur": PROJECTED_UAH_PER_EUR,
        "rateDate": SNAPSHOT_DATE.isoformat(),
    }


def project_amounts(row: dict) -> list[dict]:
    mappings = [
        ("budget", "Y", "AE"),
        ("planned_works", "Z", "AF"),
        ("planned_supervision", "AA", "AG"),
        ("planned_engineer", "AB", "AH"),
        ("eib_financing", "AC", "AI"),
        ("local_financing", "AD", "AJ"),
        ("construction", "AV", "AW"),
        ("supervision", "CA", "CB"),
        ("engineer", "CY", "CZ"),
    ]
    return [amount for kind, uah, eur in mappings if (amount := amount_record(kind, row.get(uah), row.get(eur)))]


def rounded_long(value) -> int | None:
    amount = positive_decimal(value)
    return int(amount.quantize(Decimal("1"), rounding=ROUND_HALF_UP)) if amount is not None else None


def project_tuple(row: dict, geocoding: dict[str, dict]) -> list:
    location = geocoding.get(str(row["F"]), {})
    return [
        row["uuid"], "subproject", int(decimal_value(row["B"]) or 8),
        f"(SELECT id FROM projects WHERE uuid = {sql(PROGRAMME_UUID)})",
        row["I"], row["E"], row["F"], None,
        None, row["C"], row["M"], location.get("latitude"), location.get("longitude"), row["status"], row["sector"], row["constructionType"],
        rounded_long(row["Y"]) or 0, rounded_long(row["CY"]), rounded_long(row["CA"]), rounded_long(row["AV"]),
        row["BB"], None, row["AS"], row["BC"], None, None, None, row["AS"], row["BB"], row["BC"],
        "UAH", row["AP"], None, None, None, row["AN"], row["BU"], row["BS"], row["BX"], row["CD"], row["CE"],
        row["CS"], row["CQ"], row["CV"], row["DB"], row["DC"], None,
        "(SELECT id FROM users WHERE username = 'admin' LIMIT 1)",
    ]


PROJECT_COLUMNS = [
    "uuid", "project_type", "tranche_number", "parent_project_id", "name", "site_name", "site_number", "description",
    "address", "region", "city", "latitude", "longitude", "status", "sector", "construction_type", "budget_planned",
    "engineer_consultant_contract_amount", "technical_supervision_amount", "subproject_contract_amount", "start_date", "end_date",
    "contract_signed_date", "planned_end_date", "design_contract_signing_date", "design_start_date", "design_planned_end_date",
    "construction_contract_signing_date", "construction_start_date", "projected_completion_time", "currency", "contractor_name",
    "designer_name", "design_contract_number", "design_contract_term", "construction_contract_number", "technical_supervision_name",
    "technical_supervision_contract_number", "technical_supervision_contract_date", "technical_supervision_start_date",
    "technical_supervision_planned_end_date", "engineer_consultant_name", "engineer_consultant_contract_number",
    "engineer_consultant_contract_date", "engineer_consultant_start_date", "engineer_consultant_planned_end_date", "manager_id", "created_by",
]


def expression_sql(value) -> str:
    if isinstance(value, str) and value.startswith("(SELECT "):
        return value
    return sql(value)


def build_sql(rows: list[dict], supplements: dict[str, dict[str, dict]], geocoding: dict[str, dict]) -> str:
    lines = ["-- Generated by tools/generate_urp3_seed.py. Do not edit row data manually."]
    root_values = [
        PROGRAMME_UUID, "project", 8, None, "Ukraine Recovery Programme III", "URP-III", "URP-III", None,
        None, None, None, None, None, "active", None, None, 100000000,
        None, None, None, None, None, None, None, None, None, None, None, None, None,
        "EUR", None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None,
        "(SELECT id FROM users WHERE username = 'admin' LIMIT 1)",
    ]
    lines.append("INSERT INTO projects (" + ", ".join(PROJECT_COLUMNS) + ") VALUES")
    all_values = [root_values] + [project_tuple(row, geocoding) for row in rows]
    assert all(len(values) == len(PROJECT_COLUMNS) for values in all_values), "Project seed column/value mismatch"
    lines.append(",\n".join("(" + ", ".join(expression_sql(value) for value in values) + ")" for values in all_values) + ";")

    lines.append("\nINSERT INTO programme_details (project_id, implementor, financing_institution, finance_contract_number, serapis_number, agreement_date, loan_amount, loan_currency, source_workbook, source_snapshot_date) VALUES")
    lines.append("((SELECT id FROM projects WHERE uuid = " + sql(PROGRAMME_UUID) + "), " + ", ".join(sql(value) for value in [
        "Ministry for Development of Communities and Territories of Ukraine", "European Investment Bank", "97043",
        "2023-0227", "2024-06-11", Decimal("100000000.00"), "EUR", "DB_Monitoring table.xlsx; Tranche A + B selected.xlsx", SNAPSHOT_DATE.isoformat(),
    ]) + ");")

    metadata_columns = [
        "project_id", "source_batch_id", "source_subproject_id", "source_lot_id", "name_en", "oblast_code",
        "municipality_name_uk", "municipality_name_en", "settlement_name_en", "priority_area_source",
        "project_manager_name_uk", "project_manager_name_en", "project_manager_org_id", "beneficiary_name_uk",
        "beneficiary_name_en", "beneficiary_org_id", "dream_project_id", "dream_project_url", "application_id",
        "dream_application_id", "construction_procurement_status", "construction_work_status", "geocode_accuracy",
        "geocode_query", "geocode_display_name", "source_rows", "source_workbook",
    ]
    lines.append("\nINSERT INTO project_monitoring_details (" + ", ".join(metadata_columns) + ") VALUES")
    metadata_rows = []
    for row in rows:
        metadata_rows.append([
            f"(SELECT id FROM projects WHERE uuid = {sql(row['uuid'])})", int(decimal_value(row["B"]) or 8), row["E"], row["F"],
            row["J"], row["D"], row["K"], row["L"], row["N"], row["O"], row["Q"], row["R"], str(row["S"]) if row["S"] is not None else None,
            row["T"], row["U"], str(row["V"]) if row["V"] is not None else None, row["G"], row["H"], row["W"], row["X"],
            row["AM"], row["BK"], geocoding.get(str(row["F"]), {}).get("matchLevel"),
            geocoding.get(str(row["F"]), {}).get("query"), geocoding.get(str(row["F"]), {}).get("displayName"),
            ",".join(map(str, row["sourceRows"])), "DB_Monitoring table.xlsx",
        ])
    assert all(len(values) == len(metadata_columns) for values in metadata_rows), "Monitoring seed column/value mismatch"
    lines.append(",\n".join("(" + ", ".join(expression_sql(value) for value in values) + ")" for values in metadata_rows) + ";")

    amount_rows = []
    programme_amount = {
        "kind": "budget", "amount": Decimal("100000000.00"), "currency": "EUR",
        "convertedAmount": Decimal("4800000000.00"), "uahPerEur": PROJECTED_UAH_PER_EUR, "rateDate": SNAPSHOT_DATE.isoformat(),
    }
    for project_uuid, amounts in [(PROGRAMME_UUID, [programme_amount])] + [(row["uuid"], project_amounts(row)) for row in rows]:
        for amount in amounts:
            amount_rows.append([
                f"(SELECT id FROM projects WHERE uuid = {sql(project_uuid)})", amount["kind"], amount["amount"], amount["currency"],
                amount["convertedAmount"], amount["uahPerEur"], amount["rateDate"], False,
            ])
    lines.append("\nINSERT INTO project_amounts (project_id, amount_kind, amount, currency, converted_amount, uah_per_eur, rate_date, conversion_edited) VALUES")
    lines.append(",\n".join("(" + ", ".join(expression_sql(value) for value in values) + ")" for values in amount_rows) + ";")

    procurement_rows = []
    for number, row in enumerate(rows, start=1):
        supplemental = supplements.get(str(row["E"]), {}).get("W", {})
        procurement_rows.append([
            number, int(decimal_value(row["B"]) or 8), row["C"], row["D"], row["E"], row["F"], row["AM"], row["AN"], row["AO"],
            row["AP"], row["AQ"], str(row["AR"]) if row["AR"] is not None else None, row["AS"], row["AT"],
            int(decimal_value(row["AU"])) if decimal_value(row["AU"]) is not None else None, positive_decimal(row["AV"]), positive_decimal(row["AW"]),
            decimal_value(row["AX"]), supplemental.get("procurementMethod"), supplemental.get("tenderDocumentType"),
            supplemental.get("estimatedProzorroDate"), supplemental.get("estimatedBidSubmissionDate"),
            supplemental.get("estimatedContractDate"), supplemental.get("estimatedContractEndDate"), supplemental.get("comments"), "works",
            f"(SELECT id FROM projects WHERE uuid = {sql(row['uuid'])})",
        ])
    procurement_columns = [
        "record_number", "batch_id", "oblast_name", "oblast_id", "sub_project_id", "sub_project_lot_id", "purchase_status", "tender_id",
        "prozorro_tender_id", "contractor_name_ukr", "contractor_name_eng", "contractor_id", "contract_date", "contract_end_date",
        "contract_duration_months", "contract_amount_uah", "contract_amount_eur", "financing_contract_difference_pct", "procurement_method",
        "tender_document_type", "estimated_prozorro_date", "estimated_bid_submission_date", "estimated_contract_date", "estimated_contract_end_date",
        "comments", "contract_type", "project_id",
    ]
    assert all(len(values) == len(procurement_columns) for values in procurement_rows), "Procurement seed column/value mismatch"
    lines.append("\nINSERT INTO procurement_records (" + ", ".join(procurement_columns) + ") VALUES")
    lines.append(",\n".join("(" + ", ".join(expression_sql(value) for value in values) + ")" for values in procurement_rows) + ";")
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--primary", type=Path, required=True)
    parser.add_argument("--supplement", type=Path, required=True)
    parser.add_argument("--json-output", type=Path, required=True)
    parser.add_argument("--sql-output", type=Path, required=True)
    parser.add_argument("--migration-output", type=Path)
    parser.add_argument("--geocoding", type=Path)
    args = parser.parse_args()

    rows, merge_decisions = merge_duplicate_lots(workbook_rows(args.primary))
    supplements = supplement_contracts(args.supplement)
    geocoding_document = json.loads(args.geocoding.read_text(encoding="utf-8")) if args.geocoding else {"results": {}}
    geocoding = geocoding_document.get("results", {})
    canonical = {
        "programme": {
            "uuid": PROGRAMME_UUID,
            "name": "Ukraine Recovery Programme III",
            "implementor": "Ministry for Development of Communities and Territories of Ukraine",
            "financingInstitution": "European Investment Bank",
            "financeContractNumber": "97043",
            "serapisNumber": "2023-0227",
            "agreementDate": "2024-06-11",
            "loanAmount": "100000000.00",
            "loanCurrency": "EUR",
        },
        "provenance": {
            "authoritativeWorkbook": args.primary.name,
            "supplementaryWorkbook": args.supplement.name,
            "sourceRows": 136,
            "logicalSubprojects": len(rows),
            "mergeDecisions": merge_decisions,
            "projectedUahPerEur": str(PROJECTED_UAH_PER_EUR),
            "sourceSnapshotDate": SNAPSHOT_DATE.isoformat(),
            "geocodingProvider": geocoding_document.get("provider"),
            "geocodingAttribution": geocoding_document.get("attribution"),
            "geocodingResolved": sum(1 for item in geocoding.values() if item.get("status") == "resolved"),
        },
        "subprojects": rows,
    }
    args.json_output.parent.mkdir(parents=True, exist_ok=True)
    args.sql_output.parent.mkdir(parents=True, exist_ok=True)
    args.json_output.write_text(json.dumps(canonical, ensure_ascii=False, indent=2, default=json_value), encoding="utf-8")
    seed_sql = build_sql(rows, supplements, geocoding)
    args.sql_output.write_text(seed_sql, encoding="utf-8")
    if args.migration_output:
        migration_header = """-- Rebuild OMS operational data from the authoritative URP III monitoring workbook.
-- Users, credentials, roles and login-related data are intentionally preserved.

ALTER TABLE projects
    MODIFY COLUMN address VARCHAR(500) NULL,
    MODIFY COLUMN region VARCHAR(100) NULL,
    MODIFY COLUMN city VARCHAR(100) NULL,
    MODIFY COLUMN latitude DECIMAL(10,7) NULL,
    MODIFY COLUMN longitude DECIMAL(10,7) NULL,
    MODIFY COLUMN sector VARCHAR(100) NULL,
    MODIFY COLUMN construction_type VARCHAR(100) NULL;

CREATE TABLE IF NOT EXISTS programme_details (
    project_id BIGINT NOT NULL PRIMARY KEY,
    implementor VARCHAR(500) NOT NULL,
    financing_institution VARCHAR(255) NOT NULL,
    finance_contract_number VARCHAR(100) NOT NULL,
    serapis_number VARCHAR(100) NOT NULL,
    agreement_date DATE NOT NULL,
    loan_amount DECIMAL(18,2) NOT NULL,
    loan_currency CHAR(3) NOT NULL,
    source_workbook VARCHAR(500) NOT NULL,
    source_snapshot_date DATE NULL,
    CONSTRAINT fk_programme_details_project FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE CASCADE
);

CREATE TABLE IF NOT EXISTS project_monitoring_details (
    project_id BIGINT NOT NULL PRIMARY KEY,
    source_batch_id INT NULL,
    source_subproject_id VARCHAR(100) NOT NULL,
    source_lot_id VARCHAR(100) NOT NULL,
    name_en TEXT NULL,
    oblast_code VARCHAR(32) NULL,
    municipality_name_uk VARCHAR(500) NULL,
    municipality_name_en VARCHAR(500) NULL,
    settlement_name_en VARCHAR(255) NULL,
    priority_area_source VARCHAR(255) NULL,
    project_manager_name_uk VARCHAR(500) NULL,
    project_manager_name_en VARCHAR(500) NULL,
    project_manager_org_id VARCHAR(64) NULL,
    beneficiary_name_uk VARCHAR(500) NULL,
    beneficiary_name_en VARCHAR(500) NULL,
    beneficiary_org_id VARCHAR(64) NULL,
    dream_project_id VARCHAR(128) NULL,
    dream_project_url VARCHAR(1000) NULL,
    application_id VARCHAR(128) NULL,
    dream_application_id VARCHAR(128) NULL,
    construction_procurement_status VARCHAR(255) NULL,
    construction_work_status VARCHAR(255) NULL,
    geocode_accuracy VARCHAR(32) NULL,
    geocode_query VARCHAR(1000) NULL,
    geocode_display_name VARCHAR(1000) NULL,
    source_rows VARCHAR(100) NOT NULL,
    source_workbook VARCHAR(500) NOT NULL,
    CONSTRAINT fk_monitoring_details_project FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE CASCADE,
    UNIQUE KEY uq_monitoring_source_lot (source_lot_id)
);

ALTER TABLE procurement_records
    ADD COLUMN procurement_method VARCHAR(255) NULL,
    ADD COLUMN tender_document_type VARCHAR(255) NULL,
    ADD COLUMN estimated_prozorro_date DATE NULL,
    ADD COLUMN estimated_bid_submission_date DATE NULL,
    ADD COLUMN estimated_contract_date DATE NULL,
    ADD COLUMN estimated_contract_end_date DATE NULL,
    ADD COLUMN comments TEXT NULL,
    ADD COLUMN contract_type VARCHAR(32) NOT NULL DEFAULT 'works',
    ADD COLUMN project_id BIGINT NULL;

ALTER TABLE procurement_records
    ADD CONSTRAINT fk_procurement_project FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE CASCADE,
    ADD INDEX idx_procurement_project (project_id),
    ADD INDEX idx_procurement_contract_type (contract_type);

DELETE FROM stored_file_blobs;
DELETE FROM inspection_photos;
DELETE FROM inspection_report_files;
DELETE FROM inspection_findings;
DELETE FROM financial_records;
DELETE FROM documents;
DELETE FROM incidents;
DELETE FROM audit_log;
DELETE FROM inspection_reports;
DELETE FROM project_amounts;
DELETE FROM procurement_records;
DELETE FROM projects;

"""
        args.migration_output.parent.mkdir(parents=True, exist_ok=True)
        args.migration_output.write_text(migration_header + seed_sql, encoding="utf-8")
    print(f"Generated {len(rows)} logical subprojects from 136 rows; merged {len(merge_decisions)} duplicate lot IDs.")


if __name__ == "__main__":
    main()
