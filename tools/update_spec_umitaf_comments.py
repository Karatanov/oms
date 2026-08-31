from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


PATH = Path("documentation/OMS_Supplementary_Specification.docx")
TITLE = "Appendix B - UMITAF UI and Data Clarifications"


def main() -> None:
    document = Document(PATH)
    changed = False
    for paragraph in document.paragraphs:
        if "otherwise the agreed map fallback is the relevant oblast centre" in paragraph.text:
            paragraph.text = paragraph.text.replace(
                "otherwise the agreed map fallback is the relevant oblast centre",
                "otherwise the settlement is used, with the relevant oblast centre reserved as a final explicit fallback",
            )
            changed = True
        if "confirmed address coordinates or the explicitly labelled oblast-centre fallback" in paragraph.text:
            paragraph.text = paragraph.text.replace(
                "confirmed address coordinates or the explicitly labelled oblast-centre fallback",
                "verified address coordinates, settlement coordinates or an explicitly labelled oblast-centre fallback",
            )
            changed = True
    if any(paragraph.text == TITLE for paragraph in document.paragraphs):
        if changed:
            document.save(PATH)
        print(f"Already present: {TITLE}")
        return

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading(TITLE, level=1)
    document.add_paragraph(
        "This appendix records the accepted UMITAF review comments for the Phase 1 MVP. "
        "It is normative and supersedes conflicting prototype UI descriptions."
    )
    sections = {
        "B.1 Navigation and interaction": [
            "The web client maintains browser history for screen navigation. Browser Back returns to the previous OMS screen without resetting the authenticated session.",
            "Selecting the already active Map navigation item resets the map to the Ukraine-wide overview.",
            "Interactive navigation items, tabs, links and icon actions use a pointer cursor and expose accessible labels or tooltips.",
        ],
        "B.2 Project details": [
            "General information begins with the customer-facing project, subproject or subproject-part name, followed by location data.",
            "Project detail sections are independently collapsible and use explicit expand/collapse icons.",
            "For records with coordinates, the location section provides an OpenStreetMap URL that can be opened or copied.",
        ],
        "B.3 Analytics and tables": [
            "Approved funding by oblast and subproject completion use vertically paged lists instead of horizontally scrolling bar collections.",
            "Procurement status includes every configured status, including zero values, and is rendered as a vertically readable list.",
            "Inspection-report and financial-record empty states fit the available width and do not show a redundant horizontal scrollbar.",
            "The procurement register supports 20, 50, 100 or all rows per page, Previous/Next navigation, sticky column headings and horizontal scrolling only for populated wide tables.",
        ],
        "B.4 Localisation": [
            "Imported bilingual monitoring fields are selected according to the active UI language. The English UI prefers English project, settlement, municipality, beneficiary, project-manager and contractor names when supplied; the Ukrainian UI prefers Ukrainian values.",
            "Procurement status values are displayed in one language only, according to the active UI language.",
        ],
        "B.5 Map data quality": [
            "Map popups use an opaque high-contrast background, stay inside the map viewport and auto-pan when necessary.",
            "Each imported subproject is geocoded once and the result is persisted. A confirmed street address is preferred; otherwise the settlement coordinate is used; only when the settlement cannot be resolved is the relevant oblast centre used as an explicit fallback.",
            "The checked-in URP III baseline contains 134 resolved subprojects and records geocode accuracy, query, provider display name and OpenStreetMap provenance for auditability.",
        ],
        "B.6 Financial terminology": [
            "Disbursement means programme financing allocated to screened projects and is not a synonym for money paid under certified acts.",
            "The act-based monthly chart is labelled Payments against certified works by month. It aggregates actual act/payment data converted to EUR using the stored exchange rate for the financial-record date.",
            "Allocated/disbursed programme funding and actual payments must remain separate datasets and visualisations; the system must not infer one from the other.",
        ],
    }
    for heading, statements in sections.items():
        document.add_heading(heading, level=2)
        for statement in statements:
            document.add_paragraph("• " + statement)
    document.save(PATH)
    print(PATH)


if __name__ == "__main__":
    main()
