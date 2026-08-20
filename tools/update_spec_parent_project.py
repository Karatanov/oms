from docx import Document
from docx.enum.text import WD_BREAK

path = r"D:\2026\OMS\OMS_Supplementary_Specification.docx"
document = Document(path)
needle = "Project Details clarification: When the selected record is a subproject, the page must show its parent project name in the header and in General Information. The parent-project relation is navigational metadata and does not expose internal database identifiers."

if not any(needle in paragraph.text for paragraph in document.paragraphs):
    heading = document.add_paragraph()
    heading.style = "Heading 2"
    heading.add_run("Project/subproject detail clarification")
    document.add_paragraph(needle)

document.save(path)
