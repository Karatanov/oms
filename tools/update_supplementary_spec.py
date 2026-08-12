from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

SOURCE = Path(r"D:\2026\OMS\OMS_Supplementary_Specification.docx")
OUTPUT = Path("documentation/OMS_Supplementary_Specification.docx")


def replace_text(document: Document, old: str, new: str) -> None:
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    return
            paragraph.text = paragraph.text.replace(old, new)
            return


document = Document(SOURCE)

replace_text(
    document,
    "PUT /projects/:uuid",
    "PATCH /projects/:uuid",
)
replace_text(
    document,
    "status is updated only through /projects/:uuid/status.",
    "status may be updated through this PATCH request by an authorised Project Manager or Administrator.",
)
replace_text(
    document,
    "Request body may include: name, siteName, siteNumber, address, region, city, latitude, longitude, sector, constructionType, budgetPlanned.",
    "Request body may include all editable business fields: name, siteName, siteNumber, description, address, region, city, latitude, longitude, status, sector, constructionType, budgetPlanned, engineerConsultantContractAmount, technicalSupervisionAmount, subprojectContractAmount, all project and construction dates, currency and contractorName.",
)
replace_text(
    document,
    "Returns a pre-signed MinIO download URL valid for 15 minutes. Response 200: { url: \"https://...\" }",
    "Returns the stored file directly as an attachment with its original UTF-8 filename. MVP files are stored on the application file system; MinIO remains a future storage adapter.",
)
replace_text(
    document,
    "GET /map/projects",
    "GET /projects/map",
)
replace_text(
    document,
    "Cluster radius: 80px",
    "Cluster radius: 58px",
)
replace_text(
    document,
    "Cluster icon: circle with count, colour reflects dominant status in cluster",
    "Cluster icon: primary-colour circle with the exact number of visible projects, subprojects and subproject parts in the cluster",
)
replace_text(
    document,
    "Maximum zoom for clustering: 14 (street level disables clustering)",
    "Clusters expand automatically as the user zooms in; overlapping markers are spiderfied at maximum zoom",
)

document.add_page_break()
document.add_heading("13 MVP Implementation Baseline (August 2026)", level=1)
document.add_paragraph(
    "This section is the English source-of-truth addendum for the delivered Phase 1 MVP. "
    "It records decisions and behaviour implemented after the original specification was drafted."
)

sections = {
    "13.1 Implemented Phase 1 modules": [
        "Authentication uses secure browser sessions. The web client also supports an explicit local “Remember me” option; it stores credentials only in that browser when the user chooses it.",
        "Dashboard KPI cards, recent audit activity, last-inspection photo slider and a vertical project-start-month chart are backed by live API data. The dashboard refreshes project, report and financial aggregates while it is open; the photo slider is positioned below the dashboard content.",
        "Project registry supports search, region and status filters, sortable data columns, CRUD, project details and a three-level hierarchy: project → subproject → subproject part.",
        "Project Registry bulk operations support multi-select and a single API operation to suspend or archive all selected projects.",
        "Construction type is a database ENUM and is selected from Reconstruction, Capital repair or New construction. It is rendered as a localised colour badge.",
        "Inspection reports support creation, submission, review/approval/rejection, findings, XLSX SIR import, source-file download, report reassignment, and JPEG/PNG photo upload with thumbnails and a main photo.",
        "The project Incidents (HSE) tab reads the OBSERVANCES ON HEALTH & SAFETY checklist directly from each uploaded SIR XLSX, including the answer and comment columns. It does not create duplicate incident-register records; if no SIR exists, the UI explicitly states that reports have not been uploaded.",
        "Financial monitoring supports acts, invoices, payments and advances; XLSX import/export; completed-work totals; project financial summary based on acts of completed works; and a live vertical monthly payments chart for payment and advance records.",
        "Document management supports file-system storage, typed upload, filtering, original-filename download and role-controlled permanent deletion.",
        "Administration provides a sortable user table, role and status colour badges, and full CRUD for user business data (identity, role, status, region, department and preferred language). Administrators can set or reset a password but cannot read an existing password or delete their own account.",
        "The Leaflet/OpenStreetMap map opens at Ukraine-wide extent, filters records locally, and clusters nearby projects, subprojects and subproject parts when zoomed out.",
        "Audit-log events are recorded for relevant authentication and data-change operations and exposed on the dashboard as recent activity.",
    ],
    "13.2 Project data model and editing": [
        "projects.project_type is ENUM('project', 'subproject', 'subproject_part') NOT NULL DEFAULT 'project'. A subproject must have a project parent; a subproject part must have a subproject parent.",
        "projects.construction_type is ENUM('reconstruction', 'capital_repair', 'new_construction') NOT NULL. UI labels are localised in Ukrainian and English.",
        "The edit form exposes all editable business data: identity and location, description, status, sector, construction type, budgets and contract amounts, contractor, ISO currency, coordinates and the complete project/design/construction date set. Technical UUID and creator-audit metadata remain immutable.",
        "Dates are entered with calendar controls and displayed as DD.MM.YYYY in the user interface while the REST API stores ISO YYYY-MM-DD values.",
    ],
    "13.3 Storage, download and deletion": [
        "The Phase 1 deployment stores uploaded documents, imported SIR source files and inspection photos on the application file system. Database tables contain metadata and storage paths; binary data is not stored in BLOB columns.",
        "Downloads stream the original file as an attachment and preserve the original UTF-8 filename. DELETE operations are hard deletes and are role-controlled; there is no recovery or version history in the MVP.",
    ],
    "13.4 Localisation and interaction standards": [
        "The web UI is localised for Ukrainian and English through the shared localisation catalogue. Tables use light surfaces, sortable headers, colour-coded statuses/types/roles and Material action icons with tooltips.",
        "Numeric fields restrict text input to the permitted numeric notation. Date fields use calendar pickers. Long creation and edit forms provide vertical scrolling.",
    ],
    "13.5 Phase 1 backlog identified by specification audit": [
        "Bulk reassignment and bundle export for projects and inspection reports remain to be delivered; bulk status change is implemented.",
        "The map currently filters already loaded data. Server-side bounding-box loading with debounced re-fetch, map geocoding search and a marker thumbnail are not yet delivered.",
        "The Financials monthly-payments drill-down chart and dashboard mini-map remain to be delivered.",
        "Notification delivery, offline synchronisation, external webhooks, advanced analytics and document versioning remain outside the current MVP scope as specified in Phase 2.",
    ],
}

for heading, items in sections.items():
    document.add_heading(heading, level=2)
    for item in items:
        document.add_paragraph(item)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
