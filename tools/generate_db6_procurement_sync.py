"""Build an auditable Flyway snapshot from DB6 Procurement Tracker.

The source workbook is maintained by the procurement team and is not part of
the application repository.  This tool converts its current tracker rows into
an idempotent database migration; the running OMS never reads a local file.
"""

from __future__ import annotations

import argparse
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


SHEET_NAME = "DB6 Procurement Tracker"
EXPECTED_COLUMNS = 37
CONTRACT_TYPES = {
    "W": "Роботи / Works",
    "TS": "Технічний нагляд / TS",
    "CSC": "Інженер-консультант / CSC",
}
STATUSES = {
    "Not Started": "Не розпочато / Not Started",
    "Tender Ongoing": "Закупівля триває / Tender Ongoing",
    "Contract award notice": "Повідомлення про намір укласти договір / Contract award notice",
    "Contract signed": "Договір укладено / Contract signed",
    "Cancelled": "Відмінено / Cancelled",
    "Contract terminated": "Договір розірвано / Contract terminated",
}
METHODS = {
    "National Competitive Bidding": "Національні конкурсні торги / National Competitive Bidding",
    "Local Shopping or Direct Contracting": "Місцевий шопінг/Прямий контракт / Local Shopping or Direct Contracting",
}


def clean(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        value = " ".join(value.replace("\r", " ").replace("\n", " ").split())
        return value or None
    if isinstance(value, float):
        return Decimal(str(value))
    return value


def iso_date(value):
    value = clean(value)
    if isinstance(value, str) and len(value) == 10:
        try:
            return date.fromisoformat(value).isoformat()
        except ValueError:
            return None
    return None


def sql(value) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''") + "'"


def english_part(value: str | None, known: dict[str, str]) -> str | None:
    if not value:
        return None
    for suffix, canonical in known.items():
        if value.endswith(suffix):
            return canonical
    return value


def records_from(source: Path) -> list[dict]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME]
    header = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
    if len(header) != EXPECTED_COLUMNS:
        raise ValueError(f"Expected {EXPECTED_COLUMNS} columns in {SHEET_NAME}, got {len(header)}")

    records: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if not row[0]:
            continue
        values = [clean(value) for value in row]
        procurement_id, source_code, batch_id, oblast_id = values[:4]
        type_code = str(values[9]).strip() if values[9] is not None else ""
        if not procurement_id or not source_code or type_code not in CONTRACT_TYPES:
            raise ValueError(f"Invalid source row: procurement={procurement_id!r}, subproject={source_code!r}, type={type_code!r}")
        if int(batch_id) != 8:
            raise ValueError(f"Unexpected batch {batch_id!r} in {procurement_id}; this import is limited to batch 8")
        raw_code = str(source_code).strip()
        parent_code = raw_code.split("#", 1)[0]
        lot_code = raw_code if "#Lot" in raw_code else None
        identity = (raw_code, type_code)
        if identity in seen:
            raise ValueError(f"Duplicate source procurement row {identity}")
        seen.add(identity)
        records.append({
            "source_procurement_id": str(procurement_id).strip(),
            "raw_code": raw_code,
            "parent_code": parent_code,
            "lot_code": lot_code,
            "batch_id": int(batch_id),
            "oblast_id": str(oblast_id).strip() if oblast_id else None,
            "promotor_name": values[4],
            "subproject_name_en": values[7],
            "source_contract_type": CONTRACT_TYPES[type_code],
            "type_code": type_code,
            "procurement_method": english_part(values[10], METHODS),
            "estimated_prozorro_date": iso_date(values[13]),
            "estimated_bid_submission_date": iso_date(values[14]),
            "estimated_contract_date": iso_date(values[15]),
            "estimated_contract_end_date": iso_date(values[16]),
            "estimated_total_uah": values[17],
            "estimated_eib_uah": values[18],
            "estimated_local_uah": values[19],
            "tender_id": values[20],
            "purchase_status": english_part(values[22], STATUSES),
            "contract_date": iso_date(values[25]),
            "prozorro_tender_id": values[26],
            "comments": values[27],
            "contract_amount_uah": values[32],
            "contract_duration_months": values[36],
        })
    return records


STAGE_COLUMNS = [
    "source_procurement_id", "raw_code", "parent_code", "lot_code", "batch_id", "oblast_id", "promotor_name",
    "subproject_name_en", "source_contract_type", "type_code", "procurement_method", "estimated_prozorro_date",
    "estimated_bid_submission_date", "estimated_contract_date", "estimated_contract_end_date", "estimated_total_uah",
    "estimated_eib_uah", "estimated_local_uah", "tender_id", "purchase_status", "contract_date", "prozorro_tender_id",
    "comments", "contract_amount_uah", "contract_duration_months",
]


def render(records: list[dict], source_name: str) -> str:
    rows = ",\n".join(
        "(" + ", ".join(sql(record[column]) for column in STAGE_COLUMNS) + ")" for record in records
    )
    return f"""-- Generated by tools/generate_db6_procurement_sync.py from {source_name}.
-- Authoritative snapshot of DB6 Procurement Tracker as supplied on 2026-10-05.
-- The source contains 319 Batch 8 rows (154 works, 154 TS, 11 CSC).
-- It intentionally replaces Batch 8 procurement records only; Batch 9 stays intact.

CREATE TEMPORARY TABLE procurement_db6_source (
    source_procurement_id VARCHAR(128) NOT NULL,
    raw_code VARCHAR(64) NOT NULL,
    parent_code VARCHAR(64) NOT NULL,
    lot_code VARCHAR(64) NULL,
    batch_id INT NOT NULL,
    oblast_id VARCHAR(32) NULL,
    promotor_name VARCHAR(500) NULL,
    subproject_name_en TEXT NULL,
    source_contract_type VARCHAR(255) NOT NULL,
    type_code VARCHAR(8) NOT NULL,
    procurement_method VARCHAR(255) NULL,
    estimated_prozorro_date DATE NULL,
    estimated_bid_submission_date DATE NULL,
    estimated_contract_date DATE NULL,
    estimated_contract_end_date DATE NULL,
    estimated_total_uah DECIMAL(18,2) NULL,
    estimated_eib_uah DECIMAL(18,2) NULL,
    estimated_local_uah DECIMAL(18,2) NULL,
    tender_id VARCHAR(128) NULL,
    purchase_status VARCHAR(255) NULL,
    contract_date DATE NULL,
    prozorro_tender_id VARCHAR(512) NULL,
    comments TEXT NULL,
    contract_amount_uah DECIMAL(18,2) NULL,
    contract_duration_months DECIMAL(10,2) NULL,
    PRIMARY KEY (raw_code, type_code)
);

INSERT INTO procurement_db6_source ({", ".join(STAGE_COLUMNS)}) VALUES
{rows};

-- Every source code is already represented by a subproject or subproject part.
-- Align the tranche value with the authoritative Batch 8 source before records
-- are relinked.  The hierarchy itself is deliberately unchanged.
UPDATE projects project
INNER JOIN (
    SELECT DISTINCT parent_code FROM procurement_db6_source
) source ON source.parent_code = project.site_number
SET project.tranche_number = 8
WHERE project.project_type = 'subproject';

UPDATE projects project
INNER JOIN procurement_db6_source source ON source.lot_code = project.site_number
SET project.tranche_number = 8
WHERE project.project_type = 'subproject_part';

-- DB6 is the authoritative daily tracker for actual Batch 8 procurement
-- data.  Replacing only that tranche removes stale plan rows while preserving
-- all Batch 9 records and every non-procurement entity.
DELETE FROM procurement_records WHERE batch_id = 8;

INSERT INTO procurement_records (
    batch_id, oblast_name, oblast_id, promotor_name, subproject_name_uk, subproject_name_en,
    sub_project_id, sub_project_lot_id, source_contract_type,
    subproject_total_cost_uah, estimated_total_uah, estimated_eib_uah, estimated_local_uah,
    purchase_status, tender_id, prozorro_tender_id, contract_date, contract_duration_months,
    contract_amount_uah, procurement_method, estimated_prozorro_date,
    estimated_bid_submission_date, estimated_contract_date, estimated_contract_end_date,
    local_financing_pct, comments, project_id
)
SELECT
    source.batch_id,
    COALESCE(NULLIF(parent.region, ''), CASE source.oblast_id
        WHEN 'CH' THEN 'Чернігівська' WHEN 'CK' THEN 'Черкаська' WHEN 'CV' THEN 'Чернівецька'
        WHEN 'DP' THEN 'Дніпропетровська' WHEN 'IF' THEN 'Івано-Франківська' WHEN 'KH' THEN 'Харківська'
        WHEN 'KM' THEN 'Хмельницька' WHEN 'KR' THEN 'Кіровоградська' WHEN 'KS' THEN 'Херсонська'
        WHEN 'KV' THEN 'Київська' WHEN 'LV' THEN 'Львівська' WHEN 'MY' THEN 'Миколаївська'
        WHEN 'OD' THEN 'Одеська' WHEN 'PL' THEN 'Полтавська' WHEN 'RV' THEN 'Рівненська'
        WHEN 'SM' THEN 'Сумська' WHEN 'TR' THEN 'Тернопільська' WHEN 'VN' THEN 'Вінницька'
        WHEN 'ZH' THEN 'Житомирська' WHEN 'ZK' THEN 'Закарпатська' WHEN 'ZP' THEN 'Запорізька'
        ELSE NULL END),
    source.oblast_id,
    source.promotor_name,
    NULLIF(parent.name, ''), source.subproject_name_en,
    source.parent_code, source.lot_code, source.source_contract_type,
    parent.budget_planned, source.estimated_total_uah, source.estimated_eib_uah, source.estimated_local_uah,
    source.purchase_status, source.tender_id, source.prozorro_tender_id,
    source.contract_date,
    CASE WHEN source.contract_duration_months IS NULL THEN NULL ELSE CAST(ROUND(source.contract_duration_months) AS SIGNED) END,
    source.contract_amount_uah, source.procurement_method, source.estimated_prozorro_date,
    source.estimated_bid_submission_date, source.estimated_contract_date, source.estimated_contract_end_date,
    CASE WHEN source.estimated_total_uah IS NULL OR source.estimated_total_uah = 0 OR source.estimated_local_uah IS NULL THEN NULL
         ELSE ROUND(source.estimated_local_uah * 100 / source.estimated_total_uah, 10) END,
    source.comments, COALESCE(part.id, parent.id)
FROM procurement_db6_source source
INNER JOIN projects parent
    ON parent.project_type = 'subproject'
   AND parent.site_number = source.parent_code
LEFT JOIN projects part
    ON part.project_type = 'subproject_part'
   AND part.site_number = source.lot_code;

-- All rows must survive the import and have an entity link.  A deliberate
-- migration failure is safer than silently presenting a partial tracker.
SET @expected_procurement_count = (SELECT COUNT(*) FROM procurement_db6_source);
SET @imported_procurement_count = (SELECT COUNT(*) FROM procurement_records WHERE batch_id = 8);
SET @missing_procurement_links = (SELECT COUNT(*) FROM procurement_records WHERE batch_id = 8 AND project_id IS NULL);
SET @procurement_check_sql = IF(
    @imported_procurement_count = @expected_procurement_count AND @missing_procurement_links = 0,
    'SELECT 1',
    'SIGNAL SQLSTATE ''45000'' SET MESSAGE_TEXT = ''DB6 procurement synchronization failed validation'''
);
PREPARE procurement_check FROM @procurement_check_sql;
EXECUTE procurement_check;
DEALLOCATE PREPARE procurement_check;

DROP TEMPORARY TABLE procurement_db6_source;
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--sql-output", type=Path, required=True)
    args = parser.parse_args()
    records = records_from(args.source)
    if len(records) != 319:
        raise ValueError(f"Expected 319 DB6 procurement rows, got {len(records)}")
    type_counts = {code: sum(row["type_code"] == code for row in records) for code in CONTRACT_TYPES}
    if type_counts != {"W": 154, "TS": 154, "CSC": 11}:
        raise ValueError(f"Unexpected type counts: {type_counts}")
    args.sql_output.write_text(render(records, args.source.name), encoding="utf-8")


if __name__ == "__main__":
    main()
