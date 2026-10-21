"""Generate deterministic SQL updates for the subproject registry source snapshot.

This tool is intentionally offline: it reads the customer-provided
``DB_Monitoring table.xlsx`` and writes SQL to stdout.  Database changes are
still applied exclusively through the Flyway migration generated from that
output.
"""

from __future__ import annotations

import argparse
import json
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float):
        return ("%.8f" % value).rstrip("0").rstrip(".")
    text = " ".join(str(value).replace("\r", " ").replace("\n", " ").split())
    return "" if text.startswith("#") else text


def sql(value: str) -> str:
    return "'" + value.replace("\\", "\\\\").replace("'", "''") + "'"


def hyperlink_or_value(formula, values, row: int, column: int) -> str:
    cell = formula.cell(row, column)
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    return clean(values.cell(row, column).value)


def source_rows(path: Path) -> list[tuple[str, dict[str, str]]]:
    formula = load_workbook(path, data_only=False, read_only=False)["DB Table"]
    values = load_workbook(path, data_only=True, read_only=False)["DB Table"]
    groups: OrderedDict[str, list[int]] = OrderedDict()
    for row in range(16, values.max_row + 1):
        code = clean(values.cell(row, 6).value)
        if code:
            groups.setdefault(code, []).append(row)

    # Master identity values come from the first populated source row. Financial
    # and delivery progress values come from the latest populated row, matching
    # the established URP III seed generator's reconciliation rule.
    result = []
    for lot_id, rows in groups.items():
        first, last = rows[0], rows[-1]
        def first_value(column: int) -> str:
            return next((clean(values.cell(row, column).value) for row in rows if clean(values.cell(row, column).value)), "")
        def last_value(column: int) -> str:
            return next((clean(values.cell(row, column).value) for row in reversed(rows) if clean(values.cell(row, column).value)), "")
        def first_link(column: int) -> str:
            return next((hyperlink_or_value(formula, values, row, column) for row in rows if hyperlink_or_value(formula, values, row, column)), "")
        def last_link(column: int) -> str:
            return next((hyperlink_or_value(formula, values, row, column) for row in reversed(rows) if hyperlink_or_value(formula, values, row, column)), "")

        dream_id = first_value(7)
        result.append((lot_id, {
            "pool": first_value(2),
            "name": first_value(9), "nameEn": first_value(10),
            # A data-only read of an Excel HYPERLINK cell returns its visible
            # text, not its target.  DREAM has a stable target pattern.
            "dreamUrl": f"https://dream.gov.ua/ua/project/DREAM-UA-{dream_id}/profile" if dream_id else "",
            "community": first_value(11), "communityEn": first_value(12),
            "priority": first_value(15), "constructionType": first_value(16),
            "implementer": first_value(17), "implementerEn": first_value(18),
            "beneficiary": first_value(20), "beneficiaryEn": first_value(21),
            "fundingTotal": last_value(25), "fundingWorks": last_value(26),
            "fundingTs": last_value(27), "fundingCsc": last_value(28),
            "eibCost": last_value(29), "localCost": last_value(30),
            "contractedTotal": last_value(37), "contractedEib": last_value(38),
            "contractedLocal": last_value(39), "approvedFundingPct": last_value(40),
            "contractor": last_value(48), "contractorEn": last_value(49),
            "contractDate": last_value(51), "contractEndDate": last_value(52),
            "contractDuration": last_value(53), "constructionAmount": last_value(54),
            "constructionEib": last_value(55), "constructionLocal": last_value(56),
            "paidConstruction": last_value(61), "dedUrl": first_link(62),
            "estimateUrl": first_link(63), "expertiseUrl": first_link(64),
            "permitUrl": first_link(68), "workStatus": last_value(78),
            "actNumber": last_value(79), "actDate": last_value(80),
            "completedWorkPct": last_value(81), "financingUsedPct": last_value(82),
            "paidFinancingPct": last_value(83), "workSummary": last_value(84),
            "workComments": last_value(85),
        }))
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--limit", type=int, default=25)
    args = parser.parse_args()
    for lot_id, values in source_rows(args.workbook)[args.offset:args.offset + args.limit]:
        source = json.dumps(values, ensure_ascii=False, separators=(",", ":"))
        print(
            "UPDATE project_monitoring_details SET registry_json = {json}, contractor_name_en = {contractor}, "
            "ded_url = {ded}, estimate_url = {estimate}, expertise_url = {expertise}, permit_url = {permit}, "
            "construction_act_number = {act_number}, construction_act_date = {act_date}, "
            "completed_work_pct = {completed}, financing_used_pct = {financing}, paid_financing_pct = {paid}, "
            "construction_work_summary = {summary}, construction_work_comments = {comments} "
            "WHERE source_lot_id = {lot};".format(
                json=sql(source), contractor=sql(values["contractorEn"]), ded=sql(values["dedUrl"]),
                estimate=sql(values["estimateUrl"]), expertise=sql(values["expertiseUrl"]), permit=sql(values["permitUrl"]),
                act_number=sql(values["actNumber"]), act_date=sql(values["actDate"]), completed=sql(values["completedWorkPct"]),
                financing=sql(values["financingUsedPct"]), paid=sql(values["paidFinancingPct"]),
                summary=sql(values["workSummary"]), comments=sql(values["workComments"]), lot=sql(lot_id)
            )
        )


if __name__ == "__main__":
    main()
