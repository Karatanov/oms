"""Generate the forward-only URP III Tranche B screening-data migration.

The input workbook is an operational source and is intentionally not committed.
This script emits a reviewed JSON snapshot plus an idempotent Flyway migration.
"""

from __future__ import annotations

import argparse
import json
import re
import uuid
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


NAMESPACE = uuid.UUID("cdad0bed-64ed-4a0b-a96f-67fdd4a1d308")
PROGRAMME_UUID = "42611a16-f360-5f8b-87a8-00680318888b"
SOURCE_DATE = "2026-09-15"
SOURCE_NAME = "DB_Screening results.xlsx"

# The agreed map fallback is the relevant regional centre, never a national
# default. Coordinates are kept with an explicit accuracy marker below.
REGIONS = {
    "Cherkasy": ("Черкаська", 49.4444, 32.0598),
    "Chernihiv": ("Чернігівська", 51.4982, 31.2893),
    "Chernivtsi": ("Чернівецька", 48.2921, 25.9358),
    "Dnipropetrovsk": ("Дніпропетровська", 48.4647, 35.0462),
    "Donetsk": ("Донецька", 48.0159, 37.8029),
    "Ivano-Frankivsk": ("Івано-Франківська", 48.9226, 24.7111),
    "Kharkiv": ("Харківська", 49.9935, 36.2304),
    "Kherson": ("Херсонська", 46.6354, 32.6169),
    "Khmelnytskyi": ("Хмельницька", 49.4229, 26.9871),
    "Kirovohrad": ("Кіровоградська", 48.5079, 32.2623),
    "Kyiv": ("Київська", 50.4501, 30.5234),
    "Lviv": ("Львівська", 49.8397, 24.0297),
    "Mykolaiv": ("Миколаївська", 46.9750, 31.9946),
    "Odesa": ("Одеська", 46.4825, 30.7233),
    "Poltava": ("Полтавська", 49.5883, 34.5514),
    "Rivne": ("Рівненська", 50.6199, 26.2516),
    "Sumy": ("Сумська", 50.9077, 33.4795),
    "Vinnytsya": ("Вінницька", 49.2331, 28.4682),
    "Volyn": ("Волинська", 50.7472, 25.3254),
    "Zakarpattya": ("Закарпатська", 48.6208, 22.2879),
    "Zaporizhzhya": ("Запорізька", 47.8388, 35.1396),
    "Zhytomyr": ("Житомирська", 50.2547, 28.6587),
}


def clean(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        value = " ".join(value.replace("\r", " ").replace("\n", " ").split())
        return value or None
    return value


def decimal(value) -> Decimal | None:
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    candidate = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return Decimal(candidate)
    except Exception:
        return None


def sql(value) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''") + "'"


def construction_type(activity_en: str | None) -> str:
    text = (activity_en or "").lower()
    if "new construction" in text:
        return "new_construction"
    if "overhaul" in text:
        return "capital_repair"
    # Restoration has no separate OMS construction enum. It is retained in
    # the source snapshot and is presented as the closest allowed category.
    return "reconstruction"


def sector(priority_en: str | None) -> str:
    text = (priority_en or "").lower()
    if "water supply" in text:
        return "water_supply"
    if "sewerage" in text:
        return "sewerage"
    if "heat supply" in text:
        return "heat_supply"
    if "health" in text:
        return "health"
    if "education" in text:
        return "education"
    return "shelter"


def serialisable(value):
    if isinstance(value, Decimal):
        return format(value, "f")
    return value


def source_rows(path: Path) -> list[dict]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    result = []
    for number, values in enumerate(sheet.iter_rows(min_row=5, values_only=True), start=5):
        values = [clean(value) for value in values]
        code = str(values[1] or "").strip()
        # The three 0_REL rows are reference ideas, not approved B-tranche
        # subprojects. Codes formatted XX09_XX are the source's B tranche.
        if not re.fullmatch(r"[A-Z]{2}09_\d+", code):
            continue
        region_en = str(values[5] or "").strip()
        region_uk, latitude, longitude = REGIONS[region_en]
        amount_uah = decimal(values[14])
        amount_eur = decimal(values[15])
        if amount_uah is None or amount_eur is None or amount_uah < 0 or amount_eur < 0:
            raise ValueError(f"{code}: missing/invalid UAH or EUR cost")
        result.append({
            "sourceRow": number,
            "uuid": str(uuid.uuid5(NAMESPACE, f"subproject:{code}")),
            "code": code,
            "nameUk": values[2], "nameEn": values[3],
            "region": region_uk, "regionEn": region_en,
            "municipalityUk": values[6], "municipalityEn": values[7],
            "settlementUk": values[8], "settlementEn": values[9],
            "priorityAreaUk": values[10], "priorityAreaEn": values[11],
            "activityUk": values[12], "activityEn": values[13],
            "costUah": amount_uah, "costEur": amount_eur,
            "submissionStatus": values[16], "undpChecklistUrl": values[17],
            "undpRiskComments": values[18], "undpConclusions": values[19],
            "umitafChecklistUrl": values[20], "umitafComments": values[21],
            "screeningColumns": {chr(65 + index): value for index, value in enumerate(values)},
            "latitude": latitude, "longitude": longitude,
            "sector": sector(values[11]), "constructionType": construction_type(values[13]),
        })
    if not result:
        raise ValueError("No Tranche B subprojects found")
    return result


def project_insert(row: dict) -> str:
    budget = row["costUah"].quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    address = ", ".join(item for item in (row["settlementUk"], row["region"]) if item)
    columns = [
        "uuid", "project_type", "tranche_number", "parent_project_id", "name", "site_name", "site_number",
        "description", "address", "region", "city", "latitude", "longitude", "status", "sector",
        "construction_type", "budget_planned", "currency", "created_by",
    ]
    values = [
        row["uuid"], "subproject", 9,
        "(SELECT id FROM projects WHERE uuid = " + sql(PROGRAMME_UUID) + " LIMIT 1)",
        row["nameUk"], row["code"], row["code"], row["activityUk"], address, row["region"],
        row["settlementUk"], row["latitude"], row["longitude"], "planned", row["sector"],
        row["constructionType"], budget, "UAH", "(SELECT id FROM users WHERE username = 'admin' LIMIT 1)",
    ]
    rendered = []
    for value in values:
        rendered.append(value if isinstance(value, str) and value.startswith("(SELECT ") else sql(value))
    return (
        "INSERT INTO projects (" + ", ".join(columns) + ")\nSELECT " + ", ".join(rendered) + "\n"
        "WHERE EXISTS (SELECT 1 FROM projects WHERE uuid = " + sql(PROGRAMME_UUID) + ")\n"
        "  AND NOT EXISTS (SELECT 1 FROM projects WHERE site_number = " + sql(row["code"]) + ");"
    )


def details_insert(row: dict) -> str:
    columns = [
        "project_id", "source_batch_id", "source_subproject_id", "source_lot_id", "name_en", "oblast_code",
        "municipality_name_uk", "municipality_name_en", "settlement_name_en", "priority_area_source",
        "construction_procurement_status", "construction_work_status", "geocode_accuracy", "geocode_query",
        "geocode_display_name", "source_rows", "source_workbook",
    ]
    query = ", ".join(item for item in (row["settlementUk"], row["region"], "Україна") if item)
    values = [
        "(SELECT id FROM projects WHERE uuid = " + sql(row["uuid"]) + " LIMIT 1)", 9, row["code"], row["code"],
        row["nameEn"], row["code"][:2], row["municipalityUk"], row["municipalityEn"], row["settlementEn"],
        row["priorityAreaEn"], row["submissionStatus"], None, "oblast_center", query, query,
        str(row["sourceRow"]), SOURCE_NAME,
    ]
    rendered = [value if isinstance(value, str) and value.startswith("(SELECT ") else sql(value) for value in values]
    return (
        "INSERT INTO project_monitoring_details (" + ", ".join(columns) + ")\nSELECT " + ", ".join(rendered) + "\n"
        "WHERE EXISTS (SELECT 1 FROM projects WHERE uuid = " + sql(row["uuid"]) + ")\n"
        "  AND NOT EXISTS (SELECT 1 FROM project_monitoring_details WHERE source_lot_id = " + sql(row["code"]) + ");"
    )


def amount_insert(row: dict) -> str:
    rate = (row["costUah"] / row["costEur"]).quantize(Decimal("0.00000001"), rounding=ROUND_HALF_UP)
    values = [
        "(SELECT id FROM projects WHERE uuid = " + sql(row["uuid"]) + " LIMIT 1)", "budget", row["costUah"], "UAH",
        row["costEur"], rate, SOURCE_DATE, False,
    ]
    rendered = [value if isinstance(value, str) and value.startswith("(SELECT ") else sql(value) for value in values]
    return (
        "INSERT INTO project_amounts (project_id, amount_kind, amount, currency, converted_amount, uah_per_eur, rate_date, conversion_edited)\n"
        "SELECT " + ", ".join(rendered) + "\n"
        "WHERE EXISTS (SELECT 1 FROM projects WHERE uuid = " + sql(row["uuid"]) + ")\n"
        "  AND NOT EXISTS (SELECT 1 FROM project_amounts amounts WHERE amounts.project_id = (SELECT id FROM projects WHERE uuid = " + sql(row["uuid"]) + " LIMIT 1) AND amounts.amount_kind = 'budget');"
    )


def build_migration(rows: list[dict]) -> str:
    chunks = [
        "-- Adds the approved URP III Tranche B screening subprojects.",
        "-- Generated by tools/generate_urp3_screening_tranche_b.py from DB_Screening results.xlsx.",
        "-- Only XX09_XX source codes are loaded; 0_REL reference rows are deliberately excluded.",
        "-- Geographic coordinates use the agreed regional-centre fallback and are marked in monitoring details.",
        "",
    ]
    for row in rows:
        chunks.extend((project_insert(row), details_insert(row), amount_insert(row), ""))
    return "\n".join(chunks)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--migration-output", type=Path, required=True)
    parser.add_argument("--json-output", type=Path, required=True)
    args = parser.parse_args()
    rows = source_rows(args.input)
    codes = {row["code"] for row in rows}
    assert len(rows) == len(codes), "The source contains duplicate Tranche B codes"
    assert all(re.fullmatch(r"[A-Z]{2}09_\d+", code) for code in codes), "Unexpected non-B code"
    assert all(row["costUah"] >= 0 and row["costEur"] >= 0 for row in rows), "Negative cost"
    args.migration_output.parent.mkdir(parents=True, exist_ok=True)
    args.json_output.parent.mkdir(parents=True, exist_ok=True)
    args.migration_output.write_text(build_migration(rows), encoding="utf-8")
    args.json_output.write_text(json.dumps({
        "programmeUuid": PROGRAMME_UUID,
        "tranche": "B",
        "sourceWorkbook": args.input.name,
        "sourceSnapshotDate": SOURCE_DATE,
        "subprojectCount": len(rows),
        "excludedReferenceRows": 3,
        "subprojects": rows,
    }, ensure_ascii=False, indent=2, default=serialisable), encoding="utf-8")
    migration = args.migration_output.read_text(encoding="utf-8")
    assert migration.count("INSERT INTO projects") == len(rows)
    assert migration.count("INSERT INTO project_monitoring_details") == len(rows)
    assert migration.count("INSERT INTO project_amounts") == len(rows)
    print(f"Generated and validated {len(rows)} Tranche B subprojects.")


if __name__ == "__main__":
    main()
