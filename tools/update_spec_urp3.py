from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.shared import Inches
from docx.oxml import OxmlElement


APPENDIX_TITLE = "Appendix A - URP III Operational Data Baseline"


def mark_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:tblHeader"))


def add_statement(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(label + ": ").bold = True
    paragraph.add_run(text)


def main() -> None:
    path = Path("documentation/OMS_Supplementary_Specification.docx")
    document = Document(path)
    if any(paragraph.text == APPENDIX_TITLE for paragraph in document.paragraphs):
        replacements = {
            "Unknown values are nullable; coordinates are never invented.":
                "Unknown values are nullable. Confirmed addresses use address coordinates; otherwise the agreed map fallback is the relevant oblast centre and the accuracy is recorded.",
            "The map must accept an empty marker set and must return only subprojects or subproject parts with genuine usable coordinates. All existing CRUD modules and the three-level hierarchy remain available.":
                "The map displays subprojects and subproject parts using confirmed address coordinates or the explicitly labelled oblast-centre fallback. All existing CRUD modules and the three-level hierarchy remain available.",
        }
        for paragraph in document.paragraphs:
            if paragraph.text in replacements:
                paragraph.text = replacements[paragraph.text]
            elif "coordinates are never invented" in paragraph.text:
                paragraph.text = (
                    "Unknown values: Unknown values are nullable. Confirmed addresses use address coordinates; "
                    "otherwise the agreed map fallback is the relevant oblast centre and the accuracy is recorded."
                )
        mark_header(document.tables[-1].rows[0])
        document.save(path)
        return

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading(APPENDIX_TITLE, level=1)
    document.add_paragraph(
        "This appendix is normative for the current Phase 1 operational dataset and supersedes "
        "prototype/demo-data assumptions elsewhere in this specification. The Ukrainian copy is "
        "not maintained as a separate source; this English specification remains authoritative."
    )

    document.add_heading("A.1 Authoritative programme", level=2)
    table = document.add_table(rows=0, cols=2)
    table.style = "Normal Table"
    table.autofit = False
    facts = [
        ("Programme", "Ukraine Recovery Programme III"),
        ("Implementor", "Ministry for Development of Communities and Territories of Ukraine"),
        ("Financing institution", "European Investment Bank"),
        ("Finance contract", "EIB FI 97043"),
        ("Serapis", "2023-0227"),
        ("Agreement date", "11 June 2024"),
        ("Loan amount", "EUR 100,000,000"),
        ("Operational hierarchy", "1 programme -> 134 subprojects -> 0 subproject parts"),
    ]
    for label, value in facts:
        cells = table.add_row().cells
        cells[0].width = Inches(2.1)
        cells[1].width = Inches(4.4)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].paragraphs[0].add_run(value)

    document.add_heading("A.2 Source and deterministic import", level=2)
    document.add_paragraph(
        "DB_Monitoring table.xlsx is the primary source. Tranche A + B selected.xlsx is a "
        "supplementary source for programme and procurement-plan facts. The checked-in generator "
        "tools/generate_urp3_seed.py produces canonical UTF-8 JSON and SQL with stable UUIDs."
    )
    add_statement(document, "Mapping", "One logical monitoring record is one OMS subproject.")
    add_statement(document, "Result", "136 physical rows produce 134 logical subprojects.")
    add_statement(document, "Deduplication", "Duplicate source rows SM08_09 and KH08_07#Lot6 are merged by stable lot ID.")
    add_statement(document, "No fabrication", "No inspections, financial transactions, documents or subproject parts are fabricated.")
    add_statement(document, "Unknown values", "Unknown values are nullable; coordinates are never invented.")

    document.add_heading("A.3 Data model alignment", level=2)
    model = document.add_table(rows=1, cols=3)
    model.style = "Normal Table"
    headers = ("Structure", "Purpose", "Key content")
    for index, header in enumerate(headers):
        model.rows[0].cells[index].paragraphs[0].add_run(header).bold = True
    mark_header(model.rows[0])
    rows = [
        ("projects", "Hierarchy and reusable OMS project concepts", "Programme/subproject type, parent, name/code, status, nullable location, contracts and dates"),
        ("programme_details", "One-to-one programme agreement facts", "Implementor, EIB/Serapis identifiers, agreement date, loan amount and provenance"),
        ("project_monitoring_details", "One-to-one source monitoring metadata", "Stable source/lot IDs, bilingual names, municipality, beneficiary, DREAM/application identifiers and source statuses"),
        ("project_amounts", "Exact monetary snapshots", "Source amount/currency, converted amount, UAH-per-EUR rate and rate date"),
        ("procurement_records", "Source-backed procurement state", "Owning subproject, contract type, method, planned tender/contract dates, comments and contract facts"),
    ]
    for values in rows:
        cells = model.add_row().cells
        for index, value in enumerate(values):
            cells[index].paragraphs[0].add_run(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    document.add_heading("A.4 Migration and preservation rule", level=2)
    document.add_paragraph(
        "Flyway V38 is an intentional operational-data replacement boundary. Before deployment, "
        "back up the database and upload storage. The migration deletes project-linked operational "
        "records and loads the canonical URP III dataset, while preserving users, password hashes, "
        "roles, account state and authentication data required for existing credentials to continue working."
    )
    document.add_paragraph(
        "The map must accept an empty marker set and must return only subprojects or subproject parts "
        "with genuine usable coordinates. All existing CRUD modules and the three-level hierarchy remain available."
    )

    document.save(path)


if __name__ == "__main__":
    main()
